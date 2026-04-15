"""営業リスト Excel (.xlsx) の読み書き / UPSERT ユーティリティ。

スキーマ (全タブ共通・11 列):
    1. 都道府県
    2. 業者名・事務所名
    3. 担当者名
    4. メールアドレス
    5. 問い合わせフォームURL
    6. HP
    7. 出典
    8. ステータス       (STEP 2 の送信結果で使用)
    9. 送信日時         (同)
    10. 送信方式        (同, mail / form)
    11. エラー          (同)

タブ:
    宅建業者 / 士業 / 相続専門業者 / 賃貸管理会社 / 保険代理店

UPSERT キー: (都道府県, 業者名・事務所名) のタブ内ユニーク。
"""

from __future__ import annotations

import logging
import os
from typing import Dict, Iterable, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

COLUMNS: List[str] = [
    "都道府県",
    "業者名・事務所名",
    "担当者名",
    "メールアドレス",
    "問い合わせフォームURL",
    "HP",
    "出典",
    "ステータス",
    "送信日時",
    "送信方式",
    "エラー",
]

TABS: List[str] = [
    "宅建業者",
    "士業",
    "相続専門業者",
    "賃貸管理会社",
    "保険代理店",
]

_UPSERT_KEYS = ("都道府県", "業者名・事務所名")


def _ensure_workbook(path: str) -> Workbook:
    """ファイルが存在すれば load、無ければ 5 タブをヘッダーのみ用意して新規作成。"""
    if os.path.exists(path):
        wb = load_workbook(path)
        # 足りないタブを補う
        for tab in TABS:
            if tab not in wb.sheetnames:
                ws = wb.create_sheet(title=tab)
                ws.append(COLUMNS)
        # デフォルトで作られる "Sheet" を消す (空の場合のみ)
        if "Sheet" in wb.sheetnames and wb["Sheet"].max_row <= 1 and "Sheet" not in TABS:
            del wb["Sheet"]
        return wb

    wb = Workbook()
    # デフォルトで作られる 1 枚目を最初のタブに割り当てる
    default_ws = wb.active
    default_ws.title = TABS[0]
    default_ws.append(COLUMNS)
    for tab in TABS[1:]:
        ws = wb.create_sheet(title=tab)
        ws.append(COLUMNS)
    return wb


def _key(row: Dict[str, str]) -> Tuple[str, str]:
    return tuple((row.get(k) or "").strip() for k in _UPSERT_KEYS)  # type: ignore[return-value]


def _row_to_list(row: Dict[str, str]) -> List[str]:
    return [(row.get(c) or "") for c in COLUMNS]


def _list_to_row(values: Iterable) -> Dict[str, str]:
    out: Dict[str, str] = {}
    vlist = list(values)
    for i, col in enumerate(COLUMNS):
        out[col] = str(vlist[i]) if i < len(vlist) and vlist[i] is not None else ""
    return out


def load_tab(path: str, tab_name: str) -> List[Dict[str, str]]:
    """指定タブを list[dict] で読み込む。ファイル/シートが無ければ空リスト。"""
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, read_only=True)
    if tab_name not in wb.sheetnames:
        return []
    ws = wb[tab_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    if not header:
        return []
    # 既存ヘッダーが新スキーマと違う場合も、既存列だけで辞書化
    header_list = [str(h) if h is not None else "" for h in header]
    out: List[Dict[str, str]] = []
    for values in rows_iter:
        if values is None:
            continue
        row: Dict[str, str] = {c: "" for c in COLUMNS}
        for i, col in enumerate(header_list):
            if col in COLUMNS and i < len(values):
                v = values[i]
                row[col] = "" if v is None else str(v)
        # すべて空ならスキップ
        if any((row[k] or "").strip() for k in COLUMNS):
            out.append(row)
    return out


def upsert_tab(
    path: str, tab_name: str, new_rows: Iterable[Dict[str, str]]
) -> Tuple[int, int]:
    """指定タブに UPSERT する。戻り値 (追加件数, 更新件数)。

    - ファイル / タブが無ければ作成
    - キー一致時は「新しい値が非空なら上書き、空なら既存値を維持」
    """
    if tab_name not in TABS:
        raise ValueError(f"未対応のタブ名: {tab_name}. 許可: {TABS}")

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = _ensure_workbook(path)
    ws = wb[tab_name]

    # 既存行を読み込む
    existing: List[Dict[str, str]] = []
    header_values = None
    for i, values in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header_values = values
            continue
        if values is None:
            continue
        row: Dict[str, str] = {c: "" for c in COLUMNS}
        if header_values:
            for j, col in enumerate(header_values):
                col_name = str(col) if col is not None else ""
                if col_name in COLUMNS and j < len(values):
                    v = values[j]
                    row[col_name] = "" if v is None else str(v)
        if any((row[k] or "").strip() for k in COLUMNS):
            existing.append(row)

    index: Dict[Tuple[str, str], int] = {_key(r): i for i, r in enumerate(existing)}
    added = 0
    updated = 0
    for new in new_rows:
        k = _key(new)
        if not k[0] or not k[1]:
            # キー空はスキップ
            continue
        if k in index:
            idx = index[k]
            prev = existing[idx]
            merged = {
                col: ((new.get(col) or "").strip() or prev.get(col) or "")
                for col in COLUMNS
            }
            if merged != prev:
                existing[idx] = merged
                updated += 1
        else:
            existing.append({col: (new.get(col) or "") for col in COLUMNS})
            index[k] = len(existing) - 1
            added += 1

    # シートを書き換える (ヘッダー + 全行を再書き込み)
    ws.delete_rows(1, ws.max_row)
    ws.append(COLUMNS)
    for row in existing:
        ws.append(_row_to_list(row))

    # 見やすいよう各列の幅を軽く広げる (初回のみの雑な調整)
    for i, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        current = ws.column_dimensions[letter].width or 0
        if current < 18:
            ws.column_dimensions[letter].width = 18

    wb.save(path)
    logger.info(
        "excel_io: upsert %s [%s] added=%d updated=%d total=%d",
        path, tab_name, added, updated, len(existing),
    )
    return added, updated
